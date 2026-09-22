#!/usr/bin/env python3
"""
READ ONLY stochastic validation harness for the Crystal Ball workbook.

- Reads cached values from the uploaded Crystal Ball workbook.
- Runs 10,000 Monte Carlo trials with monthly Normal(mean, stddev)
  assumptions for Sep-Dec 2026.
- Compares ERM-style Python simulation percentiles against Crystal Ball P5/P50/P95.
- Does NOT import Django, does NOT write database, does NOT modify workbook.

Usage:
  python corporate_risk/scripts/validate_crystal_ball_monte_carlo_v1_readonly.py \
      /tmp/KK_RISIKO_Penjualan_Demand_Agustus_2026.xlsx
"""
from __future__ import annotations

import math
import random
import sys
from pathlib import Path

from openpyxl import load_workbook

TRIALS = 10_000
SEED = 20_260_909
TOLERANCE_PCT = 0.05
MONTHS = ("Sep", "Okt", "Nov", "Des")


def percentile(values, q):
    ordered = sorted(float(v) for v in values)
    if not ordered:
        return 0.0
    if len(ordered) == 1:
        return ordered[0]
    pos = (len(ordered) - 1) * q
    lo = int(math.floor(pos))
    hi = int(math.ceil(pos))
    if lo == hi:
        return ordered[lo]
    frac = pos - lo
    return ordered[lo] + (ordered[hi] - ordered[lo]) * frac


def pct_diff(value, benchmark):
    if benchmark == 0:
        return 0.0 if value == 0 else math.inf
    return (value - benchmark) / benchmark * 100.0


def run_metric(actual_total, assumptions, trials, rng):
    totals = []
    for _ in range(trials):
        total = float(actual_total)
        for mean_value, stddev in assumptions:
            sampled = rng.gauss(float(mean_value), float(stddev))
            total += max(sampled, 0.0)
        totals.append(total)
    return totals


def print_metric(label, target, totals, benchmark):
    result = {
        "p5": percentile(totals, 0.05),
        "p50": percentile(totals, 0.50),
        "p95": percentile(totals, 0.95),
    }
    not_achieved = sum(1 for value in totals if value < target)
    probability_not = not_achieved / len(totals) * 100.0

    print(f"\n{label}")
    print("-" * 132)
    print(f"TARGET : {target:,.6f}")
    print(f"PROBABILITY NOT ACHIEVE TARGET : {probability_not:.4f}% ({not_achieved:,}/{len(totals):,})")
    print(f"{'PCTL':<8}{'ERM SIMULATION':>28}{'CRYSTAL BALL':>28}{'DIFF':>24}{'DIFF %':>16}{'STATUS':>14}")
    all_pass = True
    for key in ("p5", "p50", "p95"):
        diff = result[key] - benchmark[key]
        diff_pct = pct_diff(result[key], benchmark[key])
        status = "PASS" if abs(diff_pct) <= TOLERANCE_PCT else "REVIEW"
        all_pass = all_pass and status == "PASS"
        print(
            f"{key.upper():<8}"
            f"{result[key]:>28,.6f}"
            f"{benchmark[key]:>28,.6f}"
            f"{diff:>24,.6f}"
            f"{diff_pct:>15,.6f}%"
            f"{status:>14}"
        )
    return result, probability_not, all_pass


def main():
    if len(sys.argv) != 2:
        raise SystemExit(
            "Usage: python validate_crystal_ball_monte_carlo_v1_readonly.py <source.xlsx>"
        )

    source = Path(sys.argv[1]).expanduser().resolve()
    if not source.exists():
        raise SystemExit(f"SOURCE NOT FOUND: {source}")

    wb = load_workbook(source, data_only=True, read_only=False)
    required = {"Niaga", "CB_DATA_"}
    missing = required - set(wb.sheetnames)
    if missing:
        raise SystemExit(f"STOP: required sheets missing: {sorted(missing)}")

    niaga = wb["Niaga"]

    sales_actual = [float(niaga[f"Y{row}"].value) for row in range(89, 97)]
    revenue_actual = [float(niaga[f"AK{row}"].value) for row in range(89, 97)]

    sales_assumptions = [
        (float(niaga[f"AA{row}"].value), float(niaga[f"AC{row}"].value))
        for row in range(97, 101)
    ]
    revenue_assumptions = [
        (float(niaga[f"AM{row}"].value), float(niaga[f"AO{row}"].value))
        for row in range(97, 101)
    ]

    sales_benchmark = {
        "p95": float(niaga["Y117"].value),
        "p50": float(niaga["Y118"].value),
        "p5": float(niaga["Y119"].value),
    }
    revenue_benchmark = {
        "p95": float(niaga["AA117"].value),
        "p50": float(niaga["AA118"].value),
        "p5": float(niaga["AA119"].value),
    }
    hjr_benchmark = {
        "p95": float(niaga["AB117"].value),
        "p50": float(niaga["AB118"].value),
        "p5": float(niaga["AB119"].value),
    }

    sales_target = float(niaga["Y121"].value)
    revenue_target = float(niaga["AA121"].value)
    hjr_target = float(niaga["AB121"].value)

    print("=" * 132)
    print("READ ONLY STOCHASTIC VALIDATION — CRYSTAL BALL VS ERM MONTE CARLO V1")
    print("=" * 132)
    print(f"SOURCE    : {source}")
    print(f"TRIALS    : {TRIALS:,}")
    print(f"RNG SEED  : {SEED}")
    print(f"TOLERANCE : ±{TOLERANCE_PCT:.3f}% per percentile")
    print("DATABASE  : NO ACCESS / NO WRITE")

    print("\nA. IMPORTED MONTHLY ASSUMPTIONS")
    print("-" * 132)
    print(f"{'MONTH':<8}{'SALES MEAN':>22}{'SALES SD':>22}{'REVENUE MEAN':>28}{'REVENUE SD':>26}")
    for month, sales, revenue in zip(MONTHS, sales_assumptions, revenue_assumptions):
        print(
            f"{month:<8}{sales[0]:>22,.4f}{sales[1]:>22,.4f}"
            f"{revenue[0]:>28,.4f}{revenue[1]:>26,.4f}"
        )

    # Independent deterministic streams make validation reproducible.
    sales_rng = random.Random(SEED)
    revenue_rng = random.Random(SEED + 1)

    sales_totals = run_metric(sum(sales_actual), sales_assumptions, TRIALS, sales_rng)
    revenue_totals = run_metric(sum(revenue_actual), revenue_assumptions, TRIALS, revenue_rng)

    print("\nB. STOCHASTIC COMPARISON")
    sales_result, sales_probability_not, sales_pass = print_metric(
        "SALES / PENJUALAN",
        sales_target,
        sales_totals,
        sales_benchmark,
    )
    revenue_result, revenue_probability_not, revenue_pass = print_metric(
        "REVENUE / PENDAPATAN",
        revenue_target,
        revenue_totals,
        revenue_benchmark,
    )

    # Workbook HJR benchmark is the component-wise ratio of revenue percentile / sales percentile.
    hjr_result = {
        key: revenue_result[key] / sales_result[key]
        for key in ("p5", "p50", "p95")
    }
    print("\nHJR / RATE-RATIO")
    print("-" * 132)
    print(f"TARGET : {hjr_target:,.6f} Rp/kWh")
    print("SEMANTIC: derived percentile ratio = Revenue percentile / Sales percentile; HJR is NOT annual SUM.")
    print(f"{'PCTL':<8}{'ERM DERIVED':>20}{'CRYSTAL BALL':>20}{'DIFF':>18}{'DIFF %':>16}{'STATUS':>14}")
    hjr_pass = True
    for key in ("p5", "p50", "p95"):
        diff = hjr_result[key] - hjr_benchmark[key]
        diff_pct = pct_diff(hjr_result[key], hjr_benchmark[key])
        status = "PASS" if abs(diff_pct) <= TOLERANCE_PCT else "REVIEW"
        hjr_pass = hjr_pass and status == "PASS"
        print(
            f"{key.upper():<8}{hjr_result[key]:>20,.6f}{hjr_benchmark[key]:>20,.6f}"
            f"{diff:>18,.6f}{diff_pct:>15,.6f}%{status:>14}"
        )

    worst_case_impact = revenue_target - revenue_result["p5"]
    cb_worst_case_impact = revenue_target - revenue_benchmark["p5"]
    impact_diff_pct = pct_diff(worst_case_impact, cb_worst_case_impact)
    impact_pass = abs(impact_diff_pct) <= TOLERANCE_PCT

    print("\nC. EXECUTIVE-RISK REFERENCE")
    print("-" * 132)
    print(f"ERM worst-case revenue impact : Rp {worst_case_impact:,.4f}")
    print(f"CB  worst-case revenue impact : Rp {cb_worst_case_impact:,.4f}")
    print(f"Impact diff                  : {impact_diff_pct:,.6f}% | {'PASS' if impact_pass else 'REVIEW'}")
    print(f"Sales probability not achieve: {sales_probability_not:.4f}%")
    print(f"Revenue probability not achieve: {revenue_probability_not:.4f}%")

    overall = sales_pass and revenue_pass and hjr_pass and impact_pass
    print("\n" + "=" * 132)
    print(f"OVERALL VALIDATION : {'PASS' if overall else 'REVIEW'}")
    if overall:
        print("NEXT GATE: safe to implement imported-assumption mode in ERM LOCAL database/service layer.")
    else:
        print("STOP: review stochastic assumptions/tolerance before database or Executive Risk integration.")
    print("=" * 132)


if __name__ == "__main__":
    main()
