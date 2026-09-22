from __future__ import annotations

import hashlib
import math
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from corporate_risk.models import (
    MonteCarloMetricHistory,
    MultiMetricMonteCarloResult,
    RiskMetric,
)
from risk.models import ProfilRisikoKorporatItem


DEFAULT_PATH = "/tmp/08_KK_RISIKO_Kehandalan_KIT_MPP_Agustus_2026.xlsx"
YEAR = 2026
RISK_NO = 8

MONTH_NAMES = [
    "januari", "februari", "maret", "april", "mei", "juni",
    "juli", "agustus", "september", "oktober", "november", "desember",
]

KEYWORDS = [
    "target pendapatan mpp",
    "realisasi pendapatan mpp",
    "pendapatan realisasi",
    "pendapatan rkap",
    "rkap (pendapatan)",
    "total estimasi pendapatan 2026",
    "total akumulasi",
    "worst loss",
    "loss opportunity",
    "realisasi",
    "target",
]

SUSPICIOUS_TOKENS = [
    "CB_DATA",
    "Niaga",
    "#REF!",
    "[",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def as_float(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return float(v)
    try:
        return float(v)
    except Exception:
        return None


def fmt(v):
    if v is None:
        return "None"
    if isinstance(v, (int, float)):
        return f"{v:,.6f}" if abs(v) < 1_000_000 else f"{v:,.2f}"
    return str(v)


def normalize_text(v):
    return str(v).strip().lower() if v is not None else ""


def excel_col_letter(n):
    s = ""
    while n:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


# manage.py shell leaves "shell" in argv, so only accept an explicit xlsx arg.
xlsx_arg = None
for arg in sys.argv[1:]:
    if str(arg).lower().endswith(".xlsx"):
        xlsx_arg = arg
        break

path = Path(xlsx_arg or DEFAULT_PATH).expanduser().resolve()
if not path.exists():
    raise SystemExit(f"STOP: workbook tidak ditemukan: {path}")

print("=" * 160)
print("READ ONLY AUDIT — RISK #8 GANGGUAN PEMBANGKIT MPP V4.5")
print("=" * 160)
print(f"SOURCE   : {path}")
print(f"SHA256   : {sha256(path)}")
print("DB WRITE : NONE")

# Load formulas and cached values separately.
wb_formula = load_workbook(path, data_only=False, read_only=False, keep_links=True)
wb_value = load_workbook(path, data_only=True, read_only=False, keep_links=True)

print()
print("A. ERM CURRENT CONFIGURATION")
print("-" * 160)

risk = (
    ProfilRisikoKorporatItem.objects
    .filter(summary__tahun=YEAR, no_risiko=RISK_NO)
    .first()
)
if not risk:
    print(f"RISK #{RISK_NO} tahun {YEAR}: NOT FOUND")
else:
    print(f"RISK : #{risk.no_risiko} ID={risk.id} | {risk.peristiwa_risiko}")
    metrics = RiskMetric.objects.filter(corporate_risk_item=risk).order_by("id")
    for metric in metrics:
        last = (
            MonteCarloMetricHistory.objects
            .filter(metric=metric)
            .order_by("-tanggal_data", "-id")
            .first()
        )
        count_2026 = MonteCarloMetricHistory.objects.filter(
            metric=metric, tanggal_data__year=YEAR
        ).count()
        print(
            f"METRIC ID={metric.id} | {metric.name} | UNIT={metric.unit}"
            f" | DIR={metric.direction} | AGG={metric.aggregation_type}"
            f" | TARGET={metric.effective_target_value}"
            f" | TARGET_METRIC={metric.is_target_metric}"
            f" | HIST2026={count_2026}"
            f" | LAST={last.tanggal_data if last else None}"
            f" | LAST VALUE={last.metric_value if last else None}"
        )

    latest = (
        MultiMetricMonteCarloResult.objects
        .filter(corporate_risk_item=risk)
        .order_by("-created_at", "-id")
        .first()
    )
    if latest:
        print(
            f"LATEST MC RESULT: ID={latest.id}"
            f" | baseline={latest.baseline_value}"
            f" | worst={latest.worst_case_value}"
            f" | best={latest.best_case_value}"
            f" | prob={latest.probability_not_achieve_target}"
        )
    else:
        print("LATEST MC RESULT: NONE")

print()
print("B. WORKBOOK STRUCTURE")
print("-" * 160)
print(f"SHEETS ({len(wb_formula.sheetnames)}):")
for i, name in enumerate(wb_formula.sheetnames, start=1):
    ws = wb_formula[name]
    print(f"  {i:02d}. {name} | rows={ws.max_row} cols={ws.max_column}")

external_links = getattr(wb_formula, "_external_links", []) or []
print(f"EXTERNAL LINK OBJECTS : {len(external_links)}")

cb_sheets = [s for s in wb_formula.sheetnames if "cb_data" in s.lower() or "crystal" in s.lower()]
print(f"CB/CRYSTAL SHEETS     : {cb_sheets if cb_sheets else 'NONE'}")

print()
print("C. KEY LABEL / CELL DISCOVERY")
print("-" * 160)

label_hits = []
for sname in wb_formula.sheetnames:
    wsf = wb_formula[sname]
    wsv = wb_value[sname]
    for row in wsf.iter_rows():
        for cell in row:
            txt = normalize_text(cell.value)
            if not txt:
                continue
            matched = [k for k in KEYWORDS if k in txt]
            if matched:
                vcell = wsv[cell.coordinate].value
                label_hits.append((sname, cell.coordinate, cell.value, vcell, matched))
                # show context to right
                right = []
                for c in range(cell.column + 1, min(cell.column + 17, wsf.max_column + 1)):
                    coord = f"{excel_col_letter(c)}{cell.row}"
                    fv = wsf[coord].value
                    vv = wsv[coord].value
                    if fv is not None or vv is not None:
                        right.append(f"{coord}=F:{fmt(fv)} / V:{fmt(vv)}")
                print(
                    f"{sname}!{cell.coordinate} | LABEL={cell.value!r}"
                    f" | CACHED={fmt(vcell)}"
                )
                if right:
                    print("   RIGHT:", " || ".join(right[:12]))

if not label_hits:
    print("No key labels found.")

print()
print("D. MONTH HEADER / REVENUE TABLE CANDIDATES")
print("-" * 160)

table_candidates = []
for sname in wb_formula.sheetnames:
    wsf = wb_formula[sname]
    wsv = wb_value[sname]

    for row_idx in range(1, wsf.max_row + 1):
        texts = [normalize_text(wsf.cell(row_idx, c).value) for c in range(1, wsf.max_column + 1)]

        jan_cols = [i + 1 for i, t in enumerate(texts) if t == "januari"]
        for jan_col in jan_cols:
            # check contiguous Jan-Dec names
            month_cols = []
            for offset, month in enumerate(MONTH_NAMES):
                c = jan_col + offset
                if c <= wsf.max_column and month in normalize_text(wsf.cell(row_idx, c).value):
                    month_cols.append(c)
                else:
                    break

            if len(month_cols) < 6:
                continue

            print(
                f"MONTH HEADER: {sname}!{excel_col_letter(jan_col)}{row_idx}"
                f" | contiguous_months={len(month_cols)}"
            )

            # inspect up to 35 rows below for rows with numeric monthly values
            for rr in range(row_idx + 1, min(row_idx + 36, wsf.max_row + 1)):
                row_vals = [
                    as_float(wsv.cell(rr, jan_col + i).value)
                    for i in range(min(12, wsf.max_column - jan_col + 1))
                ]
                numeric_count = sum(v is not None for v in row_vals)
                if numeric_count < 4:
                    continue

                left_text = " | ".join(
                    str(wsf.cell(rr, c).value)
                    for c in range(max(1, jan_col - 3), jan_col)
                    if wsf.cell(rr, c).value not in (None, "")
                )

                jan_aug = row_vals[:8]
                sep_dec = row_vals[8:12] if len(row_vals) >= 12 else []
                jan_aug_total = (
                    sum(v for v in jan_aug if v is not None)
                    if sum(v is not None for v in jan_aug) >= 6
                    else None
                )

                score = 0
                left_norm = left_text.lower()
                if "total" in left_norm:
                    score += 3
                if "realisasi" in left_norm:
                    score += 3
                if "target" in left_norm or "rkap" in left_norm:
                    score += 2
                if numeric_count >= 8:
                    score += 2
                if numeric_count >= 12:
                    score += 1

                table_candidates.append(
                    {
                        "sheet": sname,
                        "row": rr,
                        "header_row": row_idx,
                        "jan_col": jan_col,
                        "left": left_text,
                        "values": row_vals,
                        "jan_aug_total": jan_aug_total,
                        "score": score,
                    }
                )

# print best candidates
table_candidates.sort(key=lambda x: (-x["score"], x["sheet"], x["row"]))
for cand in table_candidates[:30]:
    vals = cand["values"]
    print(
        f"{cand['sheet']}!row{cand['row']} | score={cand['score']}"
        f" | LABEL={cand['left']!r}"
        f" | JAN-AUG SUM={fmt(cand['jan_aug_total'])}"
    )
    print(
        "   JAN-DEC:",
        " | ".join(
            f"{MONTH_NAMES[i][:3].upper()}={fmt(vals[i])}"
            for i in range(min(12, len(vals)))
        )
    )

if not table_candidates:
    print("No structured Jan-Dec table candidate found.")

print()
print("E. FORMULA / STALE-LINK AUDIT")
print("-" * 160)

formula_count = 0
external_formula_count = 0
sheet_ref_counter = Counter()
suspicious_hits = []

sheet_ref_pattern = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_ .-]+))!")

for sname in wb_formula.sheetnames:
    ws = wb_formula[sname]
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if not (isinstance(value, str) and value.startswith("=")):
                continue
            formula_count += 1

            if "[" in value and "]" in value:
                external_formula_count += 1

            for m in sheet_ref_pattern.finditer(value):
                ref_sheet = (m.group(1) or m.group(2) or "").strip()
                if ref_sheet:
                    sheet_ref_counter[ref_sheet] += 1

            if any(tok.lower() in value.lower() for tok in SUSPICIOUS_TOKENS):
                suspicious_hits.append((sname, cell.coordinate, value))

print(f"FORMULAS TOTAL             : {formula_count}")
print(f"EXTERNAL-WORKBOOK FORMULAS : {external_formula_count}")

print("TOP REFERENCED SHEETS:")
for ref, cnt in sheet_ref_counter.most_common(20):
    print(f"  {ref!r}: {cnt}")

print(f"SUSPICIOUS FORMULA HITS    : {len(suspicious_hits)}")
for hit in suspicious_hits[:50]:
    print(f"  {hit[0]}!{hit[1]} = {hit[2]}")

print()
print("F. PERCENTILE / CRYSTAL BALL BENCHMARK SEARCH")
print("-" * 160)

pct_hits = []
for sname in wb_formula.sheetnames:
    wsf = wb_formula[sname]
    wsv = wb_value[sname]
    for row in wsf.iter_rows():
        for cell in row:
            txt = normalize_text(cell.value)
            if any(k in txt for k in ("p5", "p50", "p95", "percentile", "best case", "worst case", "baseline")):
                pct_hits.append(
                    (sname, cell.coordinate, cell.value, wsv[cell.coordinate].value)
                )

if pct_hits:
    for hit in pct_hits[:100]:
        print(
            f"{hit[0]}!{hit[1]} | FORMULA/LABEL={fmt(hit[2])}"
            f" | CACHED={fmt(hit[3])}"
        )
else:
    print("NO P5/P50/P95/BEST/WORST/BASELINE labels found.")

has_cb_structure = bool(cb_sheets)
has_percentile_evidence = any(
    any(k in normalize_text(hit[2]) for k in ("p5", "p50", "p95", "percentile"))
    for hit in pct_hits
)

print()
print("G. TARGET / ACTUAL / FORECAST CONSISTENCY CANDIDATES")
print("-" * 160)

# Summarize high-score monthly rows for decision support.
for cand in table_candidates[:12]:
    vals = cand["values"]
    jan_aug = vals[:8]
    sep_dec = vals[8:12]
    jan_aug_sum = cand["jan_aug_total"]
    sep_dec_sum = (
        sum(v for v in sep_dec if v is not None)
        if sep_dec and all(v is not None for v in sep_dec)
        else None
    )
    full_year = (
        jan_aug_sum + sep_dec_sum
        if jan_aug_sum is not None and sep_dec_sum is not None
        else None
    )
    print(
        f"{cand['sheet']}!row{cand['row']} | {cand['left']!r}"
        f" | JAN-AUG={fmt(jan_aug_sum)}"
        f" | SEP-DEC={fmt(sep_dec_sum)}"
        f" | FY={fmt(full_year)}"
    )

print()
print("H. DECISION GATE")
print("-" * 160)

# Determine conservative model-readiness.
if has_cb_structure and has_percentile_evidence:
    mc_gate = "REVIEW_CB_AVAILABLE"
    reason = (
        "Workbook contains CB/percentile evidence. Need exact benchmark mapping "
        "and assumption lineage before import."
    )
else:
    mc_gate = "NO_VALIDATED_CB_BENCHMARK"
    reason = (
        "No reliable Crystal Ball percentile structure was identified. "
        "Do not manufacture Monte Carlo assumptions from stale/template formulas."
    )

print(f"MONTE CARLO GATE : {mc_gate}")
print(f"REASON           : {reason}")

print()
print("SAFE NEXT:")
print("  1. Confirm the correct Jan-Aug MPP revenue row and reconcile against DB metric.")
print("  2. Confirm official full-year RKAP target.")
print("  3. Classify Sep-Dec cells as TARGET / FORECAST / FORMULA / BLANK.")
print("  4. If no valid CB benchmark exists: import ACTUAL + TARGET only; no MC validation badge.")
print("  5. If valid forecast exists but no CB benchmark: it may be stored as planning forecast,")
print("     but must not be labeled Crystal Ball validated.")
print("  6. Any stale Dashboard / Risk #2 / Niaga / external-link formulas must be excluded.")

print()
print("=" * 160)
print("READ ONLY AUDIT COMPLETE — NO DATABASE WRITE")
print("=" * 160)
