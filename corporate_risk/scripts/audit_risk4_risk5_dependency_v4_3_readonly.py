from __future__ import annotations

import math
import random
from pathlib import Path
from openpyxl import load_workbook

Z95 = 1.6448536269514722
TRIALS = 10000
SEED = 20260909

FILES = {
    4: "/tmp/04_KK_RISIKO_Pendapatan_Off_Grid_Agustus_2026.xlsx",
    5: "/tmp/05_KK_RISIKO_Gangguan_Padam_System_Agustus_2026.xlsx",
}


def pct(a, p):
    a = sorted(float(x) for x in a)
    pos = (len(a) - 1) * p
    lo = int(math.floor(pos))
    hi = int(math.ceil(pos))
    if lo == hi:
        return a[lo]
    f = pos - lo
    return a[lo] + (a[hi] - a[lo]) * f


def normal_cdf(x):
    return 0.5 * (1 + math.erf(x / math.sqrt(2)))


def implied_rho(sds, sigma_total):
    sum_sq = sum(s*s for s in sds)
    pair_sum = sum(
        sds[i] * sds[j]
        for i in range(len(sds))
        for j in range(i + 1, len(sds))
    )
    if pair_sum == 0:
        return 0.0
    return (sigma_total**2 - sum_sq) / (2 * pair_sum)


def equicorr_sim(actual_total, means, sds, rho, floor_zero=False):
    # Xi = mean_i + sd_i * (sqrt(rho) Z0 + sqrt(1-rho) Zi)
    # Valid for rho >= 0; these use-cases are non-negative.
    rng = random.Random(SEED)
    out = []
    sr = math.sqrt(max(rho, 0.0))
    si = math.sqrt(max(1.0 - rho, 0.0))
    for _ in range(TRIALS):
        z0 = rng.gauss(0, 1)
        total = actual_total
        for mean, sd in zip(means, sds):
            x = mean + sd * (sr * z0 + si * rng.gauss(0, 1))
            if floor_zero:
                x = max(x, 0)
            total += x
        out.append(total)
    return {
        "p5": pct(out, .05),
        "p50": pct(out, .50),
        "p95": pct(out, .95),
    }


def print_header(title):
    print()
    print("=" * 148)
    print(title)
    print("=" * 148)


def analyze_risk4(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Revenue Luar Batam Data Standar"]

    actual = [float(ws[f"AE{r}"].value or 0) for r in range(29, 37)]
    means = [float(ws[f"AG{r}"].value or 0) for r in range(37, 41)]
    sds = [float(ws[f"AI{r}"].value or 0) for r in range(37, 41)]

    p5 = float(ws["AF44"].value)
    p50 = float(ws["AF45"].value)
    p95 = float(ws["AF46"].value)

    actual_total = sum(actual)
    deterministic = actual_total + sum(means)

    sigma_ind = math.sqrt(sum(s*s for s in sds))
    sigma_low = (p50 - p5) / Z95
    sigma_high = (p95 - p50) / Z95
    sigma_sym = (p95 - p5) / (2 * Z95)

    rho_low = implied_rho(sds, sigma_low)
    rho_high = implied_rho(sds, sigma_high)
    rho_sym = implied_rho(sds, sigma_sym)

    indep = {
        "p5": deterministic - Z95 * sigma_ind,
        "p50": deterministic,
        "p95": deterministic + Z95 * sigma_ind,
    }
    corr = equicorr_sim(actual_total, means, sds, rho_sym, floor_zero=False)

    print_header("RISK #4 — OFF GRID | CORRELATION / DEPENDENCY DIAGNOSTIC")
    print(f"Workbook P5/P50/P95 : {p5:,.2f} | {p50:,.2f} | {p95:,.2f}")
    print(f"Deterministic mean   : {deterministic:,.2f}")
    print(f"P50 diff             : {deterministic - p50:,.6f}")
    print()
    print("MONTHLY ASSUMPTIONS")
    for i, (m, sd) in enumerate(zip(means, sds), start=9):
        pneg = normal_cdf(-m / sd) * 100 if sd else 0
        print(
            f"  2026-{i:02d} | mean={m:,.2f} | sd={sd:,.2f} "
            f"| P(monthly value < 0)={pneg:.2f}%"
        )

    print()
    print("AGGREGATE SIGMA")
    print(f"  Independent Normal sigma : {sigma_ind:,.2f}")
    print(f"  Implied from lower tail  : {sigma_low:,.2f}")
    print(f"  Implied from upper tail  : {sigma_high:,.2f}")
    print(f"  Symmetric CB sigma       : {sigma_sym:,.2f}")

    print()
    print("IMPLIED EQUAL PAIRWISE CORRELATION")
    print(f"  rho lower-tail : {rho_low:.4f}")
    print(f"  rho upper-tail : {rho_high:.4f}")
    print(f"  rho symmetric  : {rho_sym:.4f}")

    print()
    print("MODEL COMPARISON")
    print(f"{'MODEL':<28}{'P5':>22}{'P50':>22}{'P95':>22}")
    print(f"{'Workbook':<28}{p5:>22,.2f}{p50:>22,.2f}{p95:>22,.2f}")
    print(f"{'Independent Normal':<28}{indep['p5']:>22,.2f}{indep['p50']:>22,.2f}{indep['p95']:>22,.2f}")
    print(f"{'Correlated Normal rho~'+format(rho_sym,'.3f'):<28}{corr['p5']:>22,.2f}{corr['p50']:>22,.2f}{corr['p95']:>22,.2f}")

    symmetry_gap = abs(rho_high - rho_low)
    print()
    if symmetry_gap <= 0.05:
        print("CONCLUSION: STRONG CANDIDATE — correlated Normal dapat menjelaskan tail Crystal Ball secara material.")
        print("RECOMMENDATION: V4.3 Risk #4 dapat memakai monthly Normal + correlation matrix/equicorrelation; JANGAN pakai zero-floor.")
    else:
        print("CONCLUSION: REVIEW — correlation saja belum cukup menjelaskan tail.")
    return {
        "rho_low": rho_low,
        "rho_high": rho_high,
        "rho_sym": rho_sym,
        "symmetry_gap": symmetry_gap,
    }


def analyze_risk5(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Gangguan (2)"]

    actual = [float(ws[f"B{r}"].value or 0) for r in range(29, 37)]
    means = [float(ws[f"B{r}"].value or 0) for r in range(47, 51)]
    sds = [float(ws[f"D{r}"].value or 0) for r in range(47, 51)]

    p5 = float(ws["G47"].value)
    p50 = float(ws["G48"].value)
    p95 = float(ws["G49"].value)

    actual_total = sum(actual)
    deterministic = actual_total + sum(means)

    sigma_ind = math.sqrt(sum(s*s for s in sds))
    sigma_low = (p50 - p5) / Z95
    sigma_high = (p95 - p50) / Z95
    sigma_sym = (p95 - p5) / (2 * Z95)

    rho_low = implied_rho(sds, sigma_low)
    rho_high = implied_rho(sds, sigma_high)
    rho_sym = implied_rho(sds, sigma_sym)

    indep = {
        "p5": deterministic - Z95 * sigma_ind,
        "p50": deterministic,
        "p95": deterministic + Z95 * sigma_ind,
    }
    corr = equicorr_sim(actual_total, means, sds, max(rho_sym, 0), floor_zero=False)

    print_header("RISK #5 — GANGGUAN PADAM | CORRELATION / SKEW DIAGNOSTIC")
    print(f"Workbook P5/P50/P95 : {p5:,.2f} | {p50:,.2f} | {p95:,.2f}")
    print(f"Deterministic mean   : {deterministic:,.2f}")
    print(f"P50 diff             : {deterministic - p50:,.6f}")

    print()
    print("MONTHLY ASSUMPTIONS")
    for i, (m, sd) in enumerate(zip(means, sds), start=9):
        pneg = normal_cdf(-m / sd) * 100 if sd else 0
        print(
            f"  2026-{i:02d} | mean={m:,.2f} | sd={sd:,.2f} "
            f"| P(monthly value < 0)={pneg:.2f}%"
        )

    print()
    print("AGGREGATE SIGMA")
    print(f"  Independent Normal sigma : {sigma_ind:,.2f}")
    print(f"  Implied from lower tail  : {sigma_low:,.2f}")
    print(f"  Implied from upper tail  : {sigma_high:,.2f}")
    print(f"  Symmetric CB sigma       : {sigma_sym:,.2f}")

    print()
    print("IMPLIED EQUAL PAIRWISE CORRELATION")
    print(f"  rho lower-tail : {rho_low:.4f}")
    print(f"  rho upper-tail : {rho_high:.4f}")
    print(f"  rho symmetric  : {rho_sym:.4f}")

    print()
    print("MODEL COMPARISON")
    print(f"{'MODEL':<28}{'P5':>22}{'P50':>22}{'P95':>22}")
    print(f"{'Workbook':<28}{p5:>22,.2f}{p50:>22,.2f}{p95:>22,.2f}")
    print(f"{'Independent Normal':<28}{indep['p5']:>22,.2f}{indep['p50']:>22,.2f}{indep['p95']:>22,.2f}")
    print(f"{'Correlated Normal rho~'+format(rho_sym,'.3f'):<28}{corr['p5']:>22,.2f}{corr['p50']:>22,.2f}{corr['p95']:>22,.2f}")

    gap = rho_high - rho_low
    print()
    if gap > 0.10:
        print("CONCLUSION: CORRELATION-ONLY TIDAK CUKUP.")
        print("Lower tail hampir independent Normal, tetapi upper tail memerlukan variance/skew jauh lebih besar.")
        print("RECOMMENDATION: audit jenis Crystal Ball assumption / skew / compound-event sebelum import Monte Carlo Risk #5.")
    else:
        print("CONCLUSION: correlated Normal mungkin cukup.")
    return {
        "rho_low": rho_low,
        "rho_high": rho_high,
        "rho_sym": rho_sym,
        "asymmetry_gap": gap,
    }


for no, f in FILES.items():
    if not Path(f).exists():
        raise SystemExit(f"STOP: file belum ada: {f}")

r4 = analyze_risk4(FILES[4])
r5 = analyze_risk5(FILES[5])

print_header("OVERALL DECISION")
print(
    f"RISK #4 | implied rho lower={r4['rho_low']:.4f}, upper={r4['rho_high']:.4f}, "
    f"gap={r4['symmetry_gap']:.4f}"
)
print(
    f"RISK #5 | implied rho lower={r5['rho_low']:.4f}, upper={r5['rho_high']:.4f}, "
    f"gap={r5['asymmetry_gap']:.4f}"
)
print()
print("DATABASE WRITE : NONE")
print("NEXT:")
print("  Risk #4 -> candidate V4.3 correlated Normal implementation.")
print("  Risk #5 -> inspect distribution/assumption semantics before coding.")
