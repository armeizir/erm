from __future__ import annotations

import hashlib
from datetime import date
from decimal import Decimal
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from corporate_risk.models import (
    MonteCarloForecastAssumption,
    MonteCarloMetricHistory,
    RiskMetric,
)
from masterdata.models import PeriodeLaporan
from risk.models import ProfilRisikoKorporatItem


SHEET_NAME = "Gangguan (2)"
ACTUAL_ROWS = dict(zip(range(1, 9), range(29, 37)))
FORECAST_ROWS = dict(zip(range(9, 13), range(47, 51)))
VALIDATION_TOLERANCE_PCT = Decimal("0.50")


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _d(value, label) -> Decimal:
    if value in (None, ""):
        raise CommandError(f"STOP: cell workbook kosong untuk {label}.")
    return Decimal(str(value))


def _actual_d(value) -> Decimal:
    # In the Risk #5 workbook, a blank monthly compensation cell means
    # no SLA compensation for that month (zero), not missing data.
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value))


class Command(BaseCommand):
    help = (
        "Import Risk #5 Aug-2026 Crystal Ball/SARIMA baseline assumptions. "
        "Default DRY RUN. Use --sync-history to authorize audited Jan-Aug reconciliation "
        "and --apply to write LOCAL DB. Full model remains REVIEW because upper-tail P95 "
        "is not identified by the baseline Normal model."
    )

    def add_arguments(self, parser):
        parser.add_argument("xlsx")
        parser.add_argument("--apply", action="store_true")
        parser.add_argument("--sync-history", action="store_true")
        parser.add_argument("--year", type=int, default=2026)
        parser.add_argument("--risk-no", type=int, default=5)

    def handle(self, *args, **options):
        source = Path(options["xlsx"]).expanduser().resolve()
        if not source.exists():
            raise CommandError(f"Workbook tidak ditemukan: {source}")

        year = options["year"]
        risk_no = options["risk_no"]
        do_apply = bool(options["apply"])
        sync_history = bool(options["sync_history"])
        source_hash = _sha256(source)

        wb = load_workbook(source, data_only=True, read_only=True)
        if SHEET_NAME not in wb.sheetnames:
            raise CommandError(f"Sheet '{SHEET_NAME}' tidak ditemukan.")
        ws = wb[SHEET_NAME]

        actual = {
            month: _actual_d(ws[f"B{row}"].value)
            for month, row in ACTUAL_ROWS.items()
        }
        assumptions = {}
        for month, row in FORECAST_ROWS.items():
            assumptions[month] = {
                "mean": _d(ws[f"B{row}"].value, f"mean {year}-{month:02}"),
                "p15": _d(ws[f"C{row}"].value, f"p15 {year}-{month:02}"),
                "stddev": abs(_d(ws[f"D{row}"].value, f"stddev {year}-{month:02}")),
            }

        benchmark = {
            "p5": _d(ws["G47"].value, "CB P5"),
            "p50": _d(ws["G48"].value, "CB P50"),
            "p95": _d(ws["G49"].value, "CB P95"),
            "target": _d(ws["G51"].value, "CB target"),
            "probability_risk": _d(ws["G55"].value, "CB probability") * Decimal("100"),
        }

        risk = (
            ProfilRisikoKorporatItem.objects
            .filter(summary__tahun=year, no_risiko=risk_no)
            .first()
        )
        if not risk:
            raise CommandError(f"Risk #{risk_no} tahun {year} tidak ditemukan.")

        metric = (
            RiskMetric.objects
            .filter(
                corporate_risk_item=risk,
                name="Realisasi Kompensasi SLA",
            )
            .first()
        )
        if not metric:
            raise CommandError("Metric 'Realisasi Kompensasi SLA' tidak ditemukan.")

        self.stdout.write("=" * 132)
        self.stdout.write(
            f"CRYSTAL BALL BLACKOUT/SLA IMPORT V4.4.1 — "
            f"{'APPLY LOCAL' if do_apply else 'DRY RUN'}"
        )
        self.stdout.write("=" * 132)
        self.stdout.write(f"SOURCE       : {source}")
        self.stdout.write(f"SHA256       : {source_hash}")
        self.stdout.write(f"SYNC HISTORY : {'ENABLED' if sync_history else 'DISABLED'}")
        self.stdout.write(f"RISK         : #{risk.no_risiko} ID={risk.id} {risk.peristiwa_risiko}")
        self.stdout.write(
            f"METRIC       : ID={metric.id} {metric.name} | "
            f"DIR={metric.direction} | AGG={metric.aggregation_type} | "
            f"TARGET={metric.effective_target_value}"
        )
        self.stdout.write(f"YTD JAN-AUG  : Rp {sum(actual.values()):,.4f}")
        self.stdout.write(
            f"CB BENCHMARK : P5={benchmark['p5']:,.4f} "
            f"P50={benchmark['p50']:,.4f} P95={benchmark['p95']:,.4f} "
            f"P(RISK)={benchmark['probability_risk']:.4f}%"
        )
        self.stdout.write(
            "MODEL        : independent Normal monthly assumptions; truncate_at_zero=False"
        )
        self.stdout.write(
            "VALIDATION   : BASELINE PARTIAL — P5/P50/probability validated; "
            "P95 upper-tail remains external/unidentified."
        )

        mismatches = []
        missing = []
        for month, value in actual.items():
            h = (
                MonteCarloMetricHistory.objects
                .filter(
                    metric=metric,
                    tanggal_data__year=year,
                    tanggal_data__month=month,
                )
                .order_by("-id")
                .first()
            )
            if h is None:
                missing.append((month, value))
                self.stdout.write(f"MISSING  {year}-{month:02}: WORKBOOK={value}")
                continue
            diff = abs(Decimal(str(h.metric_value)) - value)
            if diff > Decimal("0.01"):
                mismatches.append((month, Decimal(str(h.metric_value)), value))
                self.stdout.write(
                    f"MISMATCH {year}-{month:02}: DB={h.metric_value} WORKBOOK={value}"
                )

        if mismatches and not sync_history:
            raise CommandError(
                "STOP: histori existing berbeda dari workbook. Tidak ada write dilakukan. "
                "Setelah rekonsiliasi disetujui, ulangi dengan --sync-history."
            )

        if sync_history:
            self.stdout.write("-" * 132)
            self.stdout.write("SYNC PLAN")
            for month, old, new in mismatches:
                self.stdout.write(f"UPDATE {year}-{month:02}: DB={old} -> WORKBOOK={new}")
            for month, new in missing:
                self.stdout.write(f"INSERT {year}-{month:02}: WORKBOOK={new}")

        if not do_apply:
            if mismatches and sync_history:
                self.stdout.write(
                    self.style.SUCCESS(
                        "DRY RUN PASS WITH SYNC PLAN — database belum diubah."
                    )
                )
            else:
                self.stdout.write(
                    self.style.SUCCESS("DRY RUN PASS — database belum diubah.")
                )
            self.stdout.write(
                "NEXT: backup DB, lalu jalankan command yang sama dengan --apply "
                "(dan --sync-history bila ada mismatch)."
            )
            return

        if mismatches and not sync_history:
            raise CommandError("STOP: --apply membutuhkan --sync-history untuk mismatch.")

        with transaction.atomic():
            desired = {
                "unit": "Rp",
                "direction": RiskMetric.DIRECTION_INCREASE,
                "aggregation_type": RiskMetric.AGGREGATION_SUM,
                "is_target_metric": True,
                "target_value": benchmark["target"],
                "is_active": True,
            }
            changed = False
            for field, value in desired.items():
                if getattr(metric, field) != value:
                    setattr(metric, field, value)
                    changed = True
            if changed:
                metric.save()

            for month, value in actual.items():
                data_date = date(year, month, 1)
                period = (
                    PeriodeLaporan.objects
                    .filter(
                        tahun_buku__tahun=year,
                        jenis_periode="bulanan",
                        tanggal_mulai__lte=data_date,
                        tanggal_selesai__gte=data_date,
                    )
                    .first()
                )
                if not period:
                    raise CommandError(f"Periode {year}-{month:02} tidak ditemukan.")

                obj = (
                    MonteCarloMetricHistory.objects
                    .filter(metric=metric, periode=period)
                    .order_by("-id")
                    .first()
                )
                if obj is None:
                    MonteCarloMetricHistory.objects.create(
                        metric=metric,
                        periode=period,
                        tanggal_data=data_date,
                        metric_value=value,
                        target_value=benchmark["target"],
                        status=MonteCarloMetricHistory.STATUS_VERIFIED,
                        keterangan=(
                            f"Crystal Ball SLA snapshot {source.name} SHA256={source_hash}"
                        ),
                    )
                else:
                    old = Decimal(str(obj.metric_value))
                    if abs(old - value) > Decimal("0.01"):
                        if not sync_history:
                            raise CommandError(
                                f"STOP: mismatch {year}-{month:02} membutuhkan --sync-history."
                            )
                        obj.metric_value = value
                        obj.target_value = benchmark["target"]
                        obj.status = MonteCarloMetricHistory.STATUS_VERIFIED
                        obj.keterangan = (
                            f"Crystal Ball SLA reconciled snapshot {source.name} "
                            f"SHA256={source_hash} | previous_value={old}"
                        )
                        obj.save()

            MonteCarloForecastAssumption.objects.filter(
                metric=metric,
                forecast_date__year=year,
                is_active=True,
            ).exclude(source_sha256=source_hash).update(is_active=False)

            dependency_model = {
                "type": "independent",
                "truncate_at_zero": False,
                "derivation": "Risk5 baseline SARIMA/Normal; upper-tail model not identified",
            }
            diagnostic = {
                "baseline_scope": "p5_p50_probability",
                "upper_tail_status": "UNRESOLVED",
                "upper_tail_reference_p95": float(benchmark["p95"]),
                "upper_tail_model": None,
                "validation_status_expected": "REVIEW",
                "reason": (
                    "Independent Normal baseline reproduces P5/P50/probability closely, "
                    "but underestimates Crystal Ball P95 by about 6.45%."
                ),
            }

            for month, data in assumptions.items():
                MonteCarloForecastAssumption.objects.update_or_create(
                    metric=metric,
                    forecast_date=date(year, month, 1),
                    source_type=MonteCarloForecastAssumption.SOURCE_CRYSTAL_BALL,
                    source_sha256=source_hash,
                    defaults={
                        "distribution_type": "normal",
                        "mean_value": data["mean"],
                        "stddev_value": data["stddev"],
                        "p15_value": data["p15"],
                        "source_file": source.name,
                        "source_sheet": SHEET_NAME,
                        "source_metadata": {
                            "model": "Crystal Ball / SARIMA baseline imported monthly assumptions",
                            "benchmark": {k: float(v) for k, v in benchmark.items()},
                            "dependency_model": dependency_model,
                            "validation_tolerance_pct": float(VALIDATION_TOLERANCE_PCT),
                            "validation_scope": "PARTIAL_BASELINE",
                            "upper_tail_diagnostic": diagnostic,
                            "validation_note": (
                                "P5/P50/probability are baseline validation checks. "
                                "P95 is retained as an external upper-tail reference and "
                                "must remain REVIEW until a severe-event driver/model is identified."
                            ),
                        },
                        "is_active": True,
                    },
                )

        self.stdout.write(
            self.style.SUCCESS(
                "APPLY PASS — Risk #5 actual Jan-Aug + independent no-floor Sep-Dec assumptions tersimpan."
            )
        )
        self.stdout.write(
            "VALIDATION STATE : PARTIAL BASELINE; full model intentionally remains REVIEW."
        )
        self.stdout.write(
            "EXECUTIVE RISK   : imported outlook will NOT be promoted while external validation status is REVIEW."
        )
        self.stdout.write(
            "NEXT             : run V4.4 baseline validator; do not deploy as fully validated model."
        )
