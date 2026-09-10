from __future__ import annotations

import hashlib
import math
from datetime import date
from decimal import Decimal
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from corporate_risk.models import (
    MonteCarloForecastAssumption,
    MonteCarloMetricHistory,
    RiskMetric,
)
from masterdata.models import PeriodeLaporan
from risk.models import ProfilRisikoKorporatItem


ACTUAL_MONTHS = list(range(1, 9))
FORECAST_MONTHS = list(range(9, 13))
SHEET_NAME = "Revenue Luar Batam Data Standar"
Z95 = Decimal("1.6448536269514722")
VALIDATION_TOLERANCE_PCT = Decimal("0.50")


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _d(value) -> Decimal:
    if value in (None, ""):
        raise CommandError("STOP: workbook berisi cell kosong pada field model yang wajib.")
    return Decimal(str(value))


def _rho_from_sigma(sigmas, sigma_total):
    sum_sq = sum(value * value for value in sigmas)
    pair_sum = sum(
        sigmas[i] * sigmas[j]
        for i in range(len(sigmas))
        for j in range(i + 1, len(sigmas))
    )
    if pair_sum == 0:
        return Decimal("0")
    return (sigma_total * sigma_total - sum_sq) / (Decimal("2") * pair_sum)


class Command(BaseCommand):
    help = (
        "Import validated Aug-2026 Crystal Ball Total Pendapatan Off Grid model to Risk #4. "
        "Default is DRY RUN; --apply writes LOCAL database."
    )

    def add_arguments(self, parser):
        parser.add_argument("xlsx")
        parser.add_argument("--apply", action="store_true")
        parser.add_argument("--year", type=int, default=2026)
        parser.add_argument("--risk-no", type=int, default=4)

    def handle(self, *args, **options):
        source = Path(options["xlsx"]).expanduser().resolve()
        if not source.exists():
            raise CommandError(f"SOURCE NOT FOUND: {source}")

        year = int(options["year"])
        risk_no = int(options["risk_no"])
        apply_mode = bool(options["apply"])
        source_hash = _sha256(source)

        wb = load_workbook(source, data_only=True, read_only=False)
        if SHEET_NAME not in wb.sheetnames:
            raise CommandError(f"Sheet '{SHEET_NAME}' tidak ditemukan.")
        ws = wb[SHEET_NAME]

        actual = {month: _d(ws[f"AE{28 + month}"].value) for month in ACTUAL_MONTHS}
        assumptions = {
            month: {
                "mean": _d(ws[f"AG{28 + month}"].value),
                "p15": _d(ws[f"AH{28 + month}"].value),
                "stddev": _d(ws[f"AI{28 + month}"].value),
            }
            for month in FORECAST_MONTHS
        }
        benchmark = {
            "p5": _d(ws["AF44"].value),
            "p50": _d(ws["AF45"].value),
            "p95": _d(ws["AF46"].value),
            "target": _d(ws["AG44"].value),
            "probability_risk": _d(ws["AI44"].value) * Decimal("100"),
        }

        deterministic_p50 = sum(actual.values()) + sum(v["mean"] for v in assumptions.values())
        if abs(deterministic_p50 - benchmark["p50"]) > Decimal("1"):
            raise CommandError(
                f"STOP: P50 workbook tidak sama dengan actual YTD + forecast means. "
                f"MODEL={deterministic_p50} WB={benchmark['p50']}"
            )

        sigmas = [abs(v["stddev"]) for v in assumptions.values()]
        sigma_lower = (benchmark["p50"] - benchmark["p5"]) / Z95
        sigma_upper = (benchmark["p95"] - benchmark["p50"]) / Z95
        sigma_symmetric = (benchmark["p95"] - benchmark["p5"]) / (Decimal("2") * Z95)
        rho_lower = _rho_from_sigma(sigmas, sigma_lower)
        rho_upper = _rho_from_sigma(sigmas, sigma_upper)
        rho_symmetric = _rho_from_sigma(sigmas, sigma_symmetric)
        rho_gap = abs(rho_upper - rho_lower)

        if not (Decimal("0") <= rho_symmetric < Decimal("1")):
            raise CommandError(f"STOP: implied equicorrelation di luar range yang didukung: {rho_symmetric}")
        if rho_gap > Decimal("0.05"):
            raise CommandError(
                f"STOP: lower/upper implied correlation tidak konsisten (gap={rho_gap}); "
                "correlation-only model tidak layak."
            )

        risk = ProfilRisikoKorporatItem.objects.filter(summary__tahun=year, no_risiko=risk_no).first()
        if not risk:
            raise CommandError(f"Risk #{risk_no} tahun {year} tidak ditemukan.")

        existing_total = RiskMetric.objects.filter(
            corporate_risk_item=risk,
            name__iexact="Total Pendapatan Off Grid",
        ).first()
        legacy_targets = list(
            RiskMetric.objects.filter(corporate_risk_item=risk, is_target_metric=True)
            .exclude(pk=existing_total.pk if existing_total else None)
            .order_by("id")
        )

        self.stdout.write("=" * 132)
        self.stdout.write("CRYSTAL BALL OFF GRID IMPORT V4.3 — " + ("APPLY LOCAL" if apply_mode else "DRY RUN"))
        self.stdout.write("=" * 132)
        self.stdout.write(f"SOURCE : {source}")
        self.stdout.write(f"SHA256 : {source_hash}")
        self.stdout.write(f"RISK   : #{risk_no} ID={risk.id} {risk.peristiwa_risiko}")
        self.stdout.write(f"YTD TOTAL OFF GRID : Rp {sum(actual.values()):,.2f}")
        self.stdout.write(
            f"CB BENCHMARK : P5={benchmark['p5']:,.2f} | P50={benchmark['p50']:,.2f} | "
            f"P95={benchmark['p95']:,.2f} | TARGET={benchmark['target']:,.2f} | "
            f"P(RISK)={benchmark['probability_risk']:,.4f}%"
        )
        self.stdout.write(
            f"DEPENDENCY : equicorrelation rho={rho_symmetric:.6f} | "
            f"rho lower={rho_lower:.6f} | rho upper={rho_upper:.6f} | gap={rho_gap:.6f}"
        )
        self.stdout.write("ZERO FLOOR  : DISABLED — workbook monthly Normal model permits negative draws.")
        self.stdout.write(f"VALIDATION TOLERANCE : ±{VALIDATION_TOLERANCE_PCT}% for sampled Crystal Ball tail benchmark.")
        if legacy_targets:
            self.stdout.write("LEGACY TARGET METRIC(S) -> will remain active but be demoted from executive target:")
            for metric in legacy_targets:
                self.stdout.write(
                    f"  ID={metric.id} | {metric.name} | TARGET={metric.effective_target_value} | WEIGHT={metric.weight}"
                )
        self.stdout.write(
            "NEW/UPDATED EXECUTIVE TARGET : Total Pendapatan Off Grid | AGG=sum | DIR=decrease | WEIGHT=1 | TARGET_METRIC=True"
        )

        if not apply_mode:
            self.stdout.write(self.style.SUCCESS("DRY RUN PASS — database belum diubah."))
            self.stdout.write("NEXT: backup DB lalu jalankan command yang sama dengan --apply.")
            return

        with transaction.atomic():
            for metric in legacy_targets:
                metric.is_target_metric = False
                metric.save(update_fields=["is_target_metric", "updated_at"])

            total_metric, _created = RiskMetric.objects.get_or_create(
                corporate_risk_item=risk,
                name="Total Pendapatan Off Grid",
                defaults={
                    "unit": "Rp",
                    "direction": RiskMetric.DIRECTION_DECREASE,
                    "aggregation_type": RiskMetric.AGGREGATION_SUM,
                    "weight": Decimal("1"),
                    "threshold": float(benchmark["target"]),
                    "is_target_metric": True,
                    "target_value": benchmark["target"],
                    "average_selling_price": Decimal("0"),
                    "is_active": True,
                },
            )
            changed = False
            desired = {
                "unit": "Rp",
                "direction": RiskMetric.DIRECTION_DECREASE,
                "aggregation_type": RiskMetric.AGGREGATION_SUM,
                "weight": Decimal("1"),
                "is_target_metric": True,
                "target_value": benchmark["target"],
                "average_selling_price": Decimal("0"),
                "is_active": True,
            }
            for field, value in desired.items():
                if getattr(total_metric, field) != value:
                    setattr(total_metric, field, value)
                    changed = True
            if changed:
                total_metric.save()

            for month, value in actual.items():
                data_date = date(year, month, 1)
                period = PeriodeLaporan.objects.filter(
                    tahun_buku__tahun=year,
                    jenis_periode="bulanan",
                    tanggal_mulai__lte=data_date,
                    tanggal_selesai__gte=data_date,
                ).first()
                if not period:
                    raise CommandError(f"Periode bulanan {year}-{month:02} tidak ditemukan.")
                obj, created = MonteCarloMetricHistory.objects.get_or_create(
                    metric=total_metric,
                    periode=period,
                    defaults={
                        "tanggal_data": data_date,
                        "metric_value": value,
                        "target_value": benchmark["target"],
                        "status": MonteCarloMetricHistory.STATUS_VERIFIED,
                        "keterangan": f"Crystal Ball Off Grid snapshot {source.name} SHA256={source_hash}",
                    },
                )
                if not created and abs(obj.metric_value - value) > Decimal("0.01"):
                    raise CommandError(
                        f"STOP: Total Pendapatan Off Grid {year}-{month:02} existing={obj.metric_value} workbook={value}"
                    )

            dependency_model = {
                "type": "equicorrelation",
                "rho": float(rho_symmetric),
                "truncate_at_zero": False,
                "derivation": "implied_from_crystal_ball_p5_p95",
            }
            diagnostic = {
                "rho_lower": float(rho_lower),
                "rho_upper": float(rho_upper),
                "rho_symmetric": float(rho_symmetric),
                "rho_tail_gap": float(rho_gap),
                "sigma_lower": float(sigma_lower),
                "sigma_upper": float(sigma_upper),
                "sigma_symmetric": float(sigma_symmetric),
            }
            for month, data in assumptions.items():
                MonteCarloForecastAssumption.objects.update_or_create(
                    metric=total_metric,
                    forecast_date=date(year, month, 1),
                    source_type=MonteCarloForecastAssumption.SOURCE_CRYSTAL_BALL,
                    source_sha256=source_hash,
                    defaults={
                        "distribution_type": "normal",
                        "mean_value": data["mean"],
                        "stddev_value": abs(data["stddev"]),
                        "p15_value": data["p15"],
                        "source_file": source.name,
                        "source_sheet": SHEET_NAME,
                        "source_metadata": {
                            "model": "Crystal Ball / ARIMA imported monthly assumptions",
                            "benchmark": {k: float(v) for k, v in benchmark.items()},
                            "dependency_model": dependency_model,
                            "dependency_diagnostic": diagnostic,
                            "validation_tolerance_pct": float(VALIDATION_TOLERANCE_PCT),
                            "validation_note": (
                                "P50 is deterministic from actual+means. Tail benchmark is a sampled Crystal Ball output; "
                                "correlation consistency and analytical correlated Normal use ±0.50% tail tolerance."
                            ),
                        },
                        "is_active": True,
                    },
                )

        self.stdout.write(self.style.SUCCESS("APPLY PASS — Total Off Grid actual Jan-Aug + correlated Sep-Dec assumptions tersimpan."))
        self.stdout.write(f"TOTAL METRIC : ID={total_metric.id} {total_metric.name}")
        self.stdout.write(f"DEPENDENCY   : equicorrelation rho={rho_symmetric:.6f}; truncate_at_zero=False")
        self.stdout.write("NEXT: run Monte Carlo imported_assumptions 10,000 trials and validator V4.3.")
