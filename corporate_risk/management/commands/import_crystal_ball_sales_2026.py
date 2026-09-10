from __future__ import annotations

import hashlib
from datetime import date
from decimal import Decimal
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from corporate_risk.models import (
    MonteCarloForecastAssumption,
    MonteCarloMetricHistory,
    RiskMetric,
)
from masterdata.models import PeriodeLaporan
from risk.models import ProfilRisikoKorporatItem


ACTUAL_MONTHS = list(range(1, 9))
FORECAST_MONTHS = list(range(9, 13))


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _d(value) -> Decimal:
    return Decimal(str(value))


class Command(BaseCommand):
    help = "Import validated Aug-2026 Crystal Ball sales/revenue/HJR linkage model. Default is DRY RUN."

    def add_arguments(self, parser):
        parser.add_argument("xlsx")
        parser.add_argument("--apply", action="store_true", help="Write histories/metrics/assumptions to LOCAL database.")
        parser.add_argument("--year", type=int, default=2026)
        parser.add_argument("--risk-no", type=int, default=2)

    def handle(self, *args, **options):
        source = Path(options["xlsx"]).expanduser().resolve()
        if not source.exists():
            raise CommandError(f"SOURCE NOT FOUND: {source}")
        year = int(options["year"])
        risk_no = int(options["risk_no"])
        apply_mode = bool(options["apply"])
        source_hash = _sha256(source)

        wb = load_workbook(source, data_only=True, read_only=False)
        if "Niaga" not in wb.sheetnames:
            raise CommandError("Sheet Niaga tidak ditemukan.")
        ws = wb["Niaga"]

        actual_sales = {month: _d(ws[f"Y{88+month}"].value) for month in ACTUAL_MONTHS}
        actual_revenue = {month: _d(ws[f"AK{88+month}"].value) for month in ACTUAL_MONTHS}
        sales_assumptions = {
            month: {
                "mean": _d(ws[f"AA{88+month}"].value),
                "p15": _d(ws[f"AB{88+month}"].value),
                "stddev": _d(ws[f"AC{88+month}"].value),
            }
            for month in FORECAST_MONTHS
        }
        revenue_assumptions = {
            month: {
                "mean": _d(ws[f"AM{88+month}"].value),
                "p15": _d(ws[f"AN{88+month}"].value),
                "stddev": _d(ws[f"AO{88+month}"].value),
            }
            for month in FORECAST_MONTHS
        }
        benchmark = {
            "sales": {"p95": _d(ws["Y117"].value), "p50": _d(ws["Y118"].value), "p5": _d(ws["Y119"].value), "target": _d(ws["Y121"].value)},
            "revenue": {"p95": _d(ws["AA117"].value), "p50": _d(ws["AA118"].value), "p5": _d(ws["AA119"].value), "target": _d(ws["AA121"].value)},
            "hjr": {"p95": _d(ws["AB117"].value), "p50": _d(ws["AB118"].value), "p5": _d(ws["AB119"].value), "target": _d(ws["AB121"].value)},
        }

        risk = ProfilRisikoKorporatItem.objects.filter(summary__tahun=year, no_risiko=risk_no).first()
        if not risk:
            raise CommandError(f"Risk #{risk_no} tahun {year} tidak ditemukan.")
        # Select the SALES metric deterministically.  Do not use a bare
        # name__icontains query here because the companion metric
        # "Pendapatan Penjualan Tenaga Listrik" contains the same phrase.
        sales_metric = (
            RiskMetric.objects.filter(
                corporate_risk_item=risk,
                name__iexact="Penjualan Tenaga Listrik (kWh)",
            ).first()
            or RiskMetric.objects.filter(
                corporate_risk_item=risk,
                is_target_metric=True,
                unit__iexact="kWh",
                name__icontains="Penjualan Tenaga Listrik",
            ).exclude(name__icontains="Pendapatan").first()
        )
        if not sales_metric:
            raise CommandError("Metric Penjualan Tenaga Listrik tidak ditemukan pada Risk #2.")
        if "pendapatan" in (sales_metric.name or "").lower():
            raise CommandError(
                f"STOP: selector sales salah memilih metric pendapatan ID={sales_metric.id} {sales_metric.name}."
            )
        if (sales_metric.unit or "").strip().lower() != "kwh":
            raise CommandError(
                f"STOP: metric sales harus berunit kWh, ditemukan ID={sales_metric.id} UNIT={sales_metric.unit}."
            )

        hjr_risk = ProfilRisikoKorporatItem.objects.filter(summary__tahun=year, no_risiko=3).first()
        hjr_metric = None
        if hjr_risk:
            hjr_metric = RiskMetric.objects.filter(
                corporate_risk_item=hjr_risk,
                name__icontains="Harga Jual Rata-rata",
            ).first()

        self.stdout.write("=" * 120)
        self.stdout.write("CRYSTAL BALL IMPORT V2 — " + ("APPLY LOCAL" if apply_mode else "DRY RUN"))
        self.stdout.write("=" * 120)
        self.stdout.write(f"SOURCE : {source}")
        self.stdout.write(f"SHA256 : {source_hash}")
        self.stdout.write(f"RISK   : #{risk_no} ID={risk.id} {risk.peristiwa_risiko}")
        self.stdout.write(f"SALES METRIC : ID={sales_metric.id} {sales_metric.name}")
        self.stdout.write(f"YTD SALES    : {sum(actual_sales.values()):,.4f}")
        self.stdout.write(f"YTD REVENUE  : {sum(actual_revenue.values()):,.4f}")
        self.stdout.write(f"CB SALES     : P5={benchmark['sales']['p5']:,.4f} P50={benchmark['sales']['p50']:,.4f} P95={benchmark['sales']['p95']:,.4f}")
        self.stdout.write(f"CB REVENUE   : P5={benchmark['revenue']['p5']:,.4f} P50={benchmark['revenue']['p50']:,.4f} P95={benchmark['revenue']['p95']:,.4f}")
        self.stdout.write(f"CB HJR       : P5={benchmark['hjr']['p5']:,.6f} P50={benchmark['hjr']['p50']:,.6f} P95={benchmark['hjr']['p95']:,.6f}")

        # Validate existing operational history before write. Never silently replace
        # a mismatched Sales or Revenue snapshot.
        mismatches = []
        for month, expected in actual_sales.items():
            existing = MonteCarloMetricHistory.objects.filter(
                metric=sales_metric,
                tanggal_data__year=year,
                tanggal_data__month=month,
            ).first()
            if existing and existing.metric_value != expected:
                mismatches.append(("SALES", month, existing.metric_value, expected))

        existing_revenue_metric = RiskMetric.objects.filter(
            corporate_risk_item=risk,
            name__iexact="Pendapatan Penjualan Tenaga Listrik",
        ).first()
        if existing_revenue_metric:
            for month, expected in actual_revenue.items():
                existing = MonteCarloMetricHistory.objects.filter(
                    metric=existing_revenue_metric,
                    tanggal_data__year=year,
                    tanggal_data__month=month,
                ).first()
                if existing and existing.metric_value != expected:
                    mismatches.append(("REVENUE", month, existing.metric_value, expected))

        if mismatches:
            for metric_kind, month, current, expected in mismatches:
                self.stderr.write(
                    f"MISMATCH {metric_kind} {year}-{month:02}: DB={current} WORKBOOK={expected}"
                )
            raise CommandError("STOP: histori existing berbeda dari workbook. Tidak ada write dilakukan.")

        if not apply_mode:
            self.stdout.write("DRY RUN PASS — database belum diubah.")
            self.stdout.write("NEXT: jalankan command yang sama dengan --apply setelah backup DB.")
            return

        with transaction.atomic():
            sales_metric.aggregation_type = RiskMetric.AGGREGATION_SUM
            sales_metric.save(update_fields=["aggregation_type", "updated_at"])

            revenue_metric, _created = RiskMetric.objects.get_or_create(
                corporate_risk_item=risk,
                name="Pendapatan Penjualan Tenaga Listrik",
                defaults={
                    "unit": "Rp",
                    "direction": RiskMetric.DIRECTION_DECREASE,
                    "aggregation_type": RiskMetric.AGGREGATION_SUM,
                    "weight": Decimal("0"),
                    "threshold": float(benchmark["revenue"]["target"]),
                    "is_target_metric": False,
                    "target_value": benchmark["revenue"]["target"],
                    "average_selling_price": Decimal("0"),
                    "is_active": True,
                },
            )
            changed = False
            if revenue_metric.aggregation_type != RiskMetric.AGGREGATION_SUM:
                revenue_metric.aggregation_type = RiskMetric.AGGREGATION_SUM
                changed = True
            if revenue_metric.target_value != benchmark["revenue"]["target"]:
                revenue_metric.target_value = benchmark["revenue"]["target"]
                changed = True
            if changed:
                revenue_metric.save()

            if hjr_metric:
                hjr_changed = False
                if hjr_metric.aggregation_type != RiskMetric.AGGREGATION_RATIO:
                    hjr_metric.aggregation_type = RiskMetric.AGGREGATION_RATIO
                    hjr_changed = True
                if hjr_metric.ratio_numerator_metric_id != revenue_metric.id:
                    hjr_metric.ratio_numerator_metric = revenue_metric
                    hjr_changed = True
                if hjr_metric.ratio_denominator_metric_id != sales_metric.id:
                    hjr_metric.ratio_denominator_metric = sales_metric
                    hjr_changed = True
                if hjr_changed:
                    hjr_metric.save(
                        update_fields=[
                            "aggregation_type",
                            "ratio_numerator_metric",
                            "ratio_denominator_metric",
                            "updated_at",
                        ]
                    )

            for metric, actual_map, target in (
                (sales_metric, actual_sales, benchmark["sales"]["target"]),
                (revenue_metric, actual_revenue, benchmark["revenue"]["target"]),
            ):
                for month, actual in actual_map.items():
                    data_date = date(year, month, 1)
                    period = PeriodeLaporan.objects.filter(
                        tahun_buku__tahun=year,
                        jenis_periode="bulanan",
                        tanggal_mulai__lte=data_date,
                        tanggal_selesai__gte=data_date,
                    ).first()
                    if not period:
                        raise CommandError(f"Periode bulanan {year}-{month:02} tidak ditemukan.")
                    obj, created = MonteCarloMetricHistory.objects.get_or_create(
                        metric=metric,
                        periode=period,
                        defaults={
                            "tanggal_data": data_date,
                            "metric_value": actual,
                            "target_value": target,
                            "status": MonteCarloMetricHistory.STATUS_VERIFIED,
                            "keterangan": f"Crystal Ball snapshot {source.name} SHA256={source_hash}",
                        },
                    )
                    if not created and obj.metric_value != actual:
                        raise CommandError(
                            f"STOP: {metric.name} {year}-{month:02} existing={obj.metric_value} workbook={actual}"
                        )

            for metric, assumption_map, metric_key in (
                (sales_metric, sales_assumptions, "sales"),
                (revenue_metric, revenue_assumptions, "revenue"),
            ):
                for month, data in assumption_map.items():
                    MonteCarloForecastAssumption.objects.update_or_create(
                        metric=metric,
                        forecast_date=date(year, month, 1),
                        source_type=MonteCarloForecastAssumption.SOURCE_CRYSTAL_BALL,
                        source_sha256=source_hash,
                        defaults={
                            "distribution_type": "normal",
                            "mean_value": data["mean"],
                            "stddev_value": data["stddev"],
                            "p15_value": data["p15"],
                            "source_file": source.name,
                            "source_sheet": "Niaga",
                            "source_metadata": {
                                "benchmark": {k: str(v) for k, v in benchmark[metric_key].items()},
                                "workbook_benchmark": {
                                    key: {k: str(v) for k, v in values.items()}
                                    for key, values in benchmark.items()
                                },
                                "workbook_marker": "CB_DATA_!P2",
                                "validated_trials": 10000,
                                "validation_tolerance_pct": "0.05",
                            },
                            "is_active": True,
                        },
                    )

        self.stdout.write(self.style.SUCCESS("APPLY PASS — histori snapshot dan Sep-Dec assumptions tersimpan."))
        self.stdout.write(f"REVENUE METRIC : ID={revenue_metric.id} {revenue_metric.name}")
        if hjr_metric:
            self.stdout.write(
                f"HJR LINK       : ID={hjr_metric.id} {hjr_metric.name} | "
                f"AGG=ratio | NUM={revenue_metric.id} | DEN={sales_metric.id} | "
                f"TARGET ERM DIPERTAHANKAN={hjr_metric.effective_target_value}"
            )
        else:
            self.stdout.write(self.style.WARNING("HJR LINK       : Risk #3 / metric HJR tidak ditemukan; dilewati."))
        self.stdout.write("NEXT: jalankan ulang Monte Carlo imported_assumptions 10,000 trials lalu cek Executive Risk V2.")
