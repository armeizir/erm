from __future__ import annotations

import hashlib
from datetime import date
from decimal import Decimal
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from corporate_risk.models import (
    MonteCarloForecastAssumption,
    MonteCarloMetricHistory,
    RiskMetric,
)
from masterdata.models import PeriodeLaporan
from risk.models import ProfilRisikoKorporatItem


ACTUAL_MONTHS = list(range(1, 9))
FORECAST_MONTHS = list(range(9, 13))
SHEET_NAME = "Energi Primer (Deny)"


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _d(value) -> Decimal:
    if value in (None, ""):
        raise CommandError("STOP: workbook berisi cell kosong pada field model yang wajib.")
    return Decimal(str(value))


def _same(a, b, tolerance=Decimal("0.0001")):
    if a is None or b is None:
        return a is b
    return abs(Decimal(a) - Decimal(b)) <= tolerance


class Command(BaseCommand):
    help = (
        "Import Aug-2026 Crystal Ball Harga Gas model to Risk #10 as RATE. "
        "Default DRY RUN; --apply writes LOCAL database. Existing ERM target is preserved."
    )

    def add_arguments(self, parser):
        parser.add_argument("xlsx")
        parser.add_argument("--apply", action="store_true", help="Write histories/assumptions to LOCAL database.")
        parser.add_argument(
            "--sync-history",
            action="store_true",
            help=(
                "Allow reconciled existing Jan-Aug gas-price history to be updated to the "
                "latest workbook snapshot. Without this flag, mismatch is a hard STOP."
            ),
        )
        parser.add_argument("--year", type=int, default=2026)
        parser.add_argument("--risk-no", type=int, default=10)

    def handle(self, *args, **options):
        source = Path(options["xlsx"]).expanduser().resolve()
        if not source.exists():
            raise CommandError(f"SOURCE NOT FOUND: {source}")

        year = int(options["year"])
        risk_no = int(options["risk_no"])
        apply_mode = bool(options["apply"])
        sync_history = bool(options["sync_history"])
        source_hash = _sha256(source)

        wb = load_workbook(source, data_only=True, read_only=False)
        if SHEET_NAME not in wb.sheetnames:
            raise CommandError(f"Sheet '{SHEET_NAME}' tidak ditemukan.")
        ws = wb[SHEET_NAME]

        actual_price = {month: _d(ws[f"N{40 + month}"].value) for month in ACTUAL_MONTHS}
        assumptions = {
            month: {
                "mean": _d(ws[f"N{53 + month}"].value),
                "p15": _d(ws[f"O{53 + month}"].value),
                "stddev": _d(ws[f"P{53 + month}"].value),
            }
            for month in FORECAST_MONTHS
        }

        workbook_ref = {
            "reported_best_case": _d(ws["D76"].value),
            "p50": _d(ws["D77"].value),
            "reported_worst_case": _d(ws["D78"].value),
            "target": _d(ws["D80"].value),
            "probability_risk": _d(ws["D84"].value) * Decimal("100"),
        }

        risk = ProfilRisikoKorporatItem.objects.filter(summary__tahun=year, no_risiko=risk_no).first()
        if not risk:
            raise CommandError(f"Risk #{risk_no} tahun {year} tidak ditemukan.")

        metric = (
            RiskMetric.objects.filter(
                corporate_risk_item=risk,
                name__iexact="Harga Gas Tertimbang Realisasi",
            ).first()
            or RiskMetric.objects.filter(
                corporate_risk_item=risk,
                unit__iexact="USD/MMBTU",
                name__icontains="Harga Gas",
            ).first()
        )
        if not metric:
            raise CommandError("Metric Harga Gas Tertimbang Realisasi tidak ditemukan pada Risk #10.")

        target_erm = metric.effective_target_value
        if target_erm in (None, ""):
            target_erm = workbook_ref["target"]
        else:
            target_erm = Decimal(str(target_erm))

        tail_inconsistent = workbook_ref["reported_best_case"] > workbook_ref["p50"]

        self.stdout.write("=" * 132)
        self.stdout.write("CRYSTAL BALL GAS PRICE IMPORT V4.2 — " + ("APPLY LOCAL" if apply_mode else "DRY RUN"))
        self.stdout.write("=" * 132)
        self.stdout.write(f"SOURCE : {source}")
        self.stdout.write(f"SHA256 : {source_hash}")
        self.stdout.write(f"SYNC HISTORY : {'ENABLED' if sync_history else 'DISABLED'}")
        self.stdout.write(f"RISK   : #{risk_no} ID={risk.id} {risk.peristiwa_risiko}")
        self.stdout.write(
            f"PRICE METRIC : ID={metric.id} {metric.name} | CURRENT AGG={metric.aggregation_type} -> REQUIRED AGG=rate | DIR={metric.direction}"
        )
        self.stdout.write(f"ACTUAL JAN-AUG AVG : {sum(actual_price.values()) / Decimal(len(actual_price)):.6f} USD/MMBTU")
        self.stdout.write(
            f"WORKBOOK REF : Best={workbook_ref['reported_best_case']:.6f} | "
            f"P50={workbook_ref['p50']:.6f} | Worst={workbook_ref['reported_worst_case']:.6f} | "
            f"Target={workbook_ref['target']:.6f} | P(RISK)={workbook_ref['probability_risk']:.4f}%"
        )
        self.stdout.write(f"TARGET ERM   : {target_erm} (dipertahankan; workbook target hanya benchmark)")
        if tail_inconsistent:
            self.stdout.write(
                self.style.WARNING(
                    "TAIL NOTE    : workbook reported Best Case > P50; D76/D78 disimpan sebagai informational reference, bukan canonical P5/P95 validation gate."
                )
            )

        mismatches = []
        for month, expected in actual_price.items():
            existing = MonteCarloMetricHistory.objects.filter(
                metric=metric,
                tanggal_data__year=year,
                tanggal_data__month=month,
            ).order_by("-id").first()
            if existing and not _same(existing.metric_value, expected, Decimal("0.0001")):
                mismatches.append((month, existing.metric_value, expected))

        if mismatches:
            for month, current, expected in mismatches:
                self.stderr.write(f"MISMATCH PRICE {year}-{month:02}: DB={current} WORKBOOK={expected}")
            if not sync_history:
                raise CommandError(
                    "STOP: histori existing berbeda dari workbook. Tidak ada write dilakukan. "
                    "Setelah rekonsiliasi disetujui, ulangi dengan --sync-history (DRY RUN), lalu --sync-history --apply."
                )
            self.stdout.write(self.style.WARNING(f"SYNC PLAN : {len(mismatches)} existing history row akan diselaraskan ke workbook."))
            for month, current, expected in mismatches:
                self.stdout.write(f"  PRICE {year}-{month:02} | DB={current} -> WORKBOOK={expected}")
        elif sync_history:
            self.stdout.write("SYNC PLAN : tidak ada existing mismatch; missing month tetap akan diinsert saat APPLY.")

        self.stdout.write("SEMANTIC CHECK : RATE = arithmetic mean of monthly gas prices; higher is worse.")
        self.stdout.write("WORST CASE     : P95 karena direction=increase.")
        self.stdout.write("VALIDATION     : canonical gate = P50 + probability at workbook target; reported Best/Worst tails remain informational due label inconsistency.")

        if not apply_mode:
            if sync_history:
                self.stdout.write(self.style.SUCCESS("DRY RUN PASS WITH SYNC PLAN — database belum diubah."))
                self.stdout.write("NEXT: backup DB lalu jalankan command yang sama dengan --sync-history --apply.")
            else:
                self.stdout.write(self.style.SUCCESS("DRY RUN PASS — database belum diubah."))
                self.stdout.write("NEXT: backup DB lalu jalankan command yang sama dengan --apply.")
            return

        with transaction.atomic():
            changed = False
            if metric.aggregation_type != RiskMetric.AGGREGATION_RATE:
                metric.aggregation_type = RiskMetric.AGGREGATION_RATE
                changed = True
            if metric.direction != RiskMetric.DIRECTION_INCREASE:
                metric.direction = RiskMetric.DIRECTION_INCREASE
                changed = True
            if not metric.is_target_metric:
                metric.is_target_metric = True
                changed = True
            # Preserve ERM target if already configured. Only initialize when absent.
            if metric.effective_target_value in (None, ""):
                metric.target_value = target_erm
                changed = True
            if changed:
                metric.save()

            for month, actual in actual_price.items():
                data_date = date(year, month, 1)
                period = PeriodeLaporan.objects.filter(
                    tahun_buku__tahun=year,
                    jenis_periode="bulanan",
                    tanggal_mulai__lte=data_date,
                    tanggal_selesai__gte=data_date,
                ).first()
                if not period:
                    raise CommandError(f"Periode bulanan {year}-{month:02} tidak ditemukan.")

                obj, created = MonteCarloMetricHistory.objects.get_or_create(
                    metric=metric,
                    periode=period,
                    defaults={
                        "tanggal_data": data_date,
                        "metric_value": actual,
                        "target_value": target_erm,
                        "status": MonteCarloMetricHistory.STATUS_VERIFIED,
                        "keterangan": f"Crystal Ball Gas Price snapshot {source.name} SHA256={source_hash}",
                    },
                )
                if not created and not _same(obj.metric_value, actual, Decimal("0.0001")):
                    if not sync_history:
                        raise CommandError(
                            f"STOP: {metric.name} {year}-{month:02} existing={obj.metric_value} workbook={actual}"
                        )
                    previous_value = obj.metric_value
                    obj.tanggal_data = data_date
                    obj.metric_value = actual
                    obj.target_value = target_erm
                    obj.status = MonteCarloMetricHistory.STATUS_VERIFIED
                    obj.keterangan = (
                        f"Crystal Ball Gas Price reconciled snapshot {source.name} "
                        f"SHA256={source_hash}; previous_value={previous_value}"
                    )
                    obj.save(update_fields=[
                        "tanggal_data", "metric_value", "target_value", "status", "keterangan", "updated_at"
                    ])
                elif not created:
                    changed_fields = []
                    if obj.target_value != target_erm:
                        obj.target_value = target_erm
                        changed_fields.append("target_value")
                    if obj.status != MonteCarloMetricHistory.STATUS_VERIFIED:
                        obj.status = MonteCarloMetricHistory.STATUS_VERIFIED
                        changed_fields.append("status")
                    if not obj.keterangan:
                        obj.keterangan = f"Crystal Ball Gas Price snapshot {source.name} SHA256={source_hash}"
                        changed_fields.append("keterangan")
                    if changed_fields:
                        changed_fields.append("updated_at")
                        obj.save(update_fields=changed_fields)

            # Only P50 and probability are authoritative gates for this workbook.  The
            # reported Best/Worst cells are preserved verbatim as informational tails.
            benchmark = {
                "p50": str(workbook_ref["p50"]),
                "target": str(workbook_ref["target"]),
                "probability_risk": str(workbook_ref["probability_risk"]),
            }
            for month, data in assumptions.items():
                MonteCarloForecastAssumption.objects.update_or_create(
                    metric=metric,
                    forecast_date=date(year, month, 1),
                    source_type=MonteCarloForecastAssumption.SOURCE_CRYSTAL_BALL,
                    source_sha256=source_hash,
                    defaults={
                        "distribution_type": "normal",
                        "mean_value": data["mean"],
                        "stddev_value": data["stddev"],
                        "p15_value": data["p15"],
                        "source_file": source.name,
                        "source_sheet": SHEET_NAME,
                        "source_metadata": {
                            "benchmark": benchmark,
                            "workbook_reported_tails": {
                                "best_case_cell_D76": str(workbook_ref["reported_best_case"]),
                                "baseline_p50_cell_D77": str(workbook_ref["p50"]),
                                "worst_case_cell_D78": str(workbook_ref["reported_worst_case"]),
                                "tail_label_inconsistent": bool(tail_inconsistent),
                            },
                            "target_erm": str(target_erm),
                            "target_workbook_reference": str(workbook_ref["target"]),
                            "validated_trials": 10000,
                            "validation_tolerance_pct": "0.05",
                            "probability_tolerance_percentage_point": "0.10",
                            "validation_scope": "P50 + analytical probability at workbook target; reported Best/Worst are informational",
                            "risk_semantics": "higher_is_worse",
                            "aggregation_semantics": "RATE_arithmetic_mean_monthly",
                            "worst_case_percentile": "P95",
                        },
                        "is_active": True,
                    },
                )

        if sync_history:
            self.stdout.write(self.style.SUCCESS("APPLY PASS WITH HISTORY SYNC — Risk #10 RATE snapshot/assumptions tersimpan."))
        else:
            self.stdout.write(self.style.SUCCESS("APPLY PASS — Risk #10 RATE snapshot/assumptions tersimpan."))
        self.stdout.write(
            f"PRICE METRIC : ID={metric.id} | DIR={metric.direction} | AGG={metric.aggregation_type} | TARGET ERM={metric.effective_target_value}"
        )
        self.stdout.write("NEXT: run Risk #10 imported_assumptions 10,000 trials dan validator V4.2.")
